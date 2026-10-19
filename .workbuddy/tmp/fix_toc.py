# -*- coding: utf-8 -*-
"""补充目录配套修正：
1) 摘要/Abstract 移出目录（克隆 Heading 1 为无大纲级别的 FrontTitle 样式，外观不变）
2) 前置部分（封面+摘要+Abstract+目录）设大写罗马数字页码，封面不显示页码
3) 目录一级条目加粗（与参考模板一致）
"""
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

P = r"F:\基于springboot框架MoveNet模型人体姿态评估康复训练辅助平台\论文初稿加开题答辩\毕业论文初稿-基于SpringBoot与MoveNet的康复训练辅助平台.docx"

d = Document(P)

# ---------- 1) 克隆 Heading 1 -> FrontTitle（去掉 outlineLvl，不进目录） ----------
styles_el = d.styles.element
h1el = d.styles["Heading 1"]._element
# 若已存在则先删
for st in styles_el.findall(qn("w:style")):
    if st.get(qn("w:styleId")) == "FrontTitle":
        styles_el.remove(st)
new = deepcopy(h1el)
new.set(qn("w:styleId"), "FrontTitle")
nm = new.find(qn("w:name"))
nm.set(qn("w:val"), "FrontTitle")
# 去掉 outlineLvl（目录 \o 开关按大纲级别抓取）
pPr = new.find(qn("w:pPr"))
if pPr is not None:
    ol = pPr.find(qn("w:outlineLvl"))
    if ol is not None:
        pPr.remove(ol)
# link 指向原字符样式，去掉避免歧义
lk = new.find(qn("w:link"))
if lk is not None:
    new.remove(lk)
styles_el.append(new)

front_style = d.styles["FrontTitle"]
n_changed = 0
for para in d.paragraphs:
    t = para.text.strip()
    if t in ("摘　　要", "Abstract"):
        para.style = front_style
        n_changed += 1
print("移出目录的标题数:", n_changed)

# ---------- 2) 前置部分罗马数字页码 ----------
sec0 = d.sections[0]
sectPr = sec0._sectPr
pg = sectPr.find(qn("w:pgNumType"))
if pg is None:
    pg = OxmlElement("w:pgNumType")
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sectPr.append(pg)
pg.set(qn("w:fmt"), "upperRoman")
print("section0 pgNumType:", dict(pg.attrib))

# 封面首页不显示页码
sec0.different_first_page_header_footer = True

# 前置部分页脚：PAGE 域（按节格式渲染为罗马数字），右对齐
ftr = sec0.footer
ftr.is_linked_to_previous = False
fp = ftr.paragraphs[0]
for r in list(fp.runs):
    r._r.getparent().remove(r._r)
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
pf = fp.paragraph_format
pf.first_line_indent = Pt(0)
pf.line_spacing = 1.0


def add_field_run(par, instr_text, placeholder):
    rb = par.add_run()
    fb = OxmlElement("w:fldChar")
    fb.set(qn("w:fldCharType"), "begin")
    rb._r.append(fb)

    ri = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr_text
    ri._r.append(it)

    rs = par.add_run()
    fs = OxmlElement("w:fldChar")
    fs.set(qn("w:fldCharType"), "separate")
    rs._r.append(fs)

    rt = par.add_run()
    tt = OxmlElement("w:t")
    tt.text = placeholder
    rt._r.append(tt)

    re_ = par.add_run()
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    re_._r.append(fe)
    return [rb, ri, rs, rt, re_]


runs = add_field_run(fp, " PAGE ", "I")
for r in runs:
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    rPr = r._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:eastAsia"), "宋体")
print("前置页脚 PAGE 域已添加")

# ---------- 3) 目录一级条目加粗 ----------
for s in d.styles:
    try:
        sname = s.name or ""
    except Exception:
        continue
    if sname.lower() == "toc 1":
        s.font.bold = True
        print("toc 1 已加粗")

d.save(P)
print("saved OK")
