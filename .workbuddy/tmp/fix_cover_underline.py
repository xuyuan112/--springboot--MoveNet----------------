# -*- coding: utf-8 -*-
"""封面下划线优化：
- 六行标签统一为 5 个全角字符宽（题　　目：/学　　院：/专业名称：/学　　号：/学生姓名：/指导教师：），列对齐
- 短内容行：居中制表位+右制表位实现"等长下划线"，内容在下划线内居中（不用手工空格）
- 题目行：恢复为本课题题目，下划线随文字自然换行（与模板一致）
- 学号恢复为用户本人的 202530495223（当前封面残留的是参考论文谭明硕的 202420495207）
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

P = r"F:\基于springboot框架MoveNet模型人体姿态评估康复训练辅助平台\论文初稿加开题答辩\毕业论文初稿-基于SpringBoot与MoveNet的康复训练辅助平台.docx"

INDENT = Cm(3.97)      # 标签列起点（沿用原首行缩进位置）
S = 3.97 + 5 * 16 / 72 * 2.54   # 标签右缘 = 缩进 + 5个全角字宽(16pt) ≈ 6.79cm
SPAN = 4.6             # 下划线总长 cm
M = round(S + SPAN / 2, 2)      # 居中制表位
R = round(S + SPAN, 2)          # 右制表位
print(f"标签右缘 S={S:.2f}cm, 居中位 M={M}cm, 右边界 R={R}cm")

TITLE = "基于SpringBoot与MoveNet的康复训练辅助平台的设计与实现"

LINES = [
    (7,  "题　　目：", TITLE,  True),
    (9,  "学　　院：", "计算机学院", False),
    (10, "专业名称：", "软件工程", False),
    (11, "学　　号：", "202530495223", False),
    (12, "学生姓名：", "许袁", False),
    (13, "指导教师：", "冯珊", False),
]


def set_font(run, cn_font, sz, ascii_font="Times New Roman"):
    run.font.size = Pt(sz)
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    rf.set(qn("w:ascii"), ascii_font)
    rf.set(qn("w:hAnsi"), ascii_font)
    rf.set(qn("w:eastAsia"), cn_font)


d = Document(P)
for idx, label, content, is_title in LINES:
    para = d.paragraphs[idx]
    # 清空原有 runs
    for r in list(para.runs):
        r._r.getparent().remove(r._r)
    pf = para.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.left_indent = INDENT
    pf.first_line_indent = Pt(0)
    if is_title:
        # 题目行：标签 + 带下划线的题目文字，自然换行；行距与其他行一致用 1.5 倍
        pf.line_spacing = 1.5
        r1 = para.add_run(label)
        set_font(r1, "华文仿宋", 16)
        r2 = para.add_run(content)
        set_font(r2, "宋体", 16)
        r2.font.underline = True
    else:
        pf.tab_stops.clear_all()
        pf.tab_stops.add_tab_stop(Cm(M), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(R), WD_TAB_ALIGNMENT.RIGHT)
        r1 = para.add_run(label)
        set_font(r1, "华文仿宋", 16)
        r2 = para.add_run("\t" + content + "\t")
        set_font(r2, "华文仿宋", 16)
        r2.font.underline = True
    print(f"段{idx} 已重建: {label}{content}")

d.save(P)
print("saved OK")
