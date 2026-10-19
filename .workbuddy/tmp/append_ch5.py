# -*- coding: utf-8 -*-
"""向论文追加 第五章 系统实现（5.1 开发环境 / 5.2 公共功能 / 5.3 核心业务功能 / 5.4 端侧算法）
图片位置以带虚线框的占位段落标记，待后续插入真实运行截图。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"F:\基于springboot框架MoveNet模型人体姿态评估康复训练辅助平台\论文初稿加开题答辩\毕业论文初稿-基于SpringBoot与MoveNet的康复训练辅助平台.docx"

CN, CN_H, EN, MONO = "宋体", "黑体", "Times New Roman", "Consolas"


def run_font(run, cn=CN, en=EN, size=None, bold=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.name = en
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)
    rf.set(qn('w:cs'), en)


def first_line_indent_chars(p, chars=200):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), '480')


def no_indent(p):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), '0')
    ind.set(qn('w:firstLine'), '0')


def para_border(p, color='BFBFBF', sz=6, val='dashed'):
    pPr = p._p.get_or_add_pPr()
    old = pPr.find(qn('w:pBdr'))
    if old is not None:
        pPr.remove(old)
    bdr = OxmlElement('w:pBdr')
    for tag in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), val)
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '4')
        el.set(qn('w:color'), color)
        bdr.append(el)
    pPr.append(bdr)


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = Pt(22)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    first_line_indent_chars(p)
    return p


def caption(doc, text, color=None):
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(3)
    c.paragraph_format.space_after = Pt(6)
    c.paragraph_format.line_spacing = 1.0
    r = c.add_run(text)
    run_font(r, CN, EN, 10.5)
    if color is not None:
        r.font.color.rgb = color
    return c


def img_placeholder(doc, cap):
    """图片占位：虚线框 + 灰字提示，后接正式题注"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    para_border(p)
    r = p.add_run('（此处预留图片位置，待插入运行截图）')
    run_font(r, CN, EN, 10.5)
    r.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
    caption(doc, cap)


def code_block(doc, cap, lines):
    caption(doc, cap)
    for ln in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Cm(0.4)
        no_indent(p)
        r = p.add_run(ln if ln else ' ')
        run_font(r, CN, MONO, 9)
    doc.add_paragraph().paragraph_format.line_spacing = 1.0


def set_table_h_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for tag, val, sz in [('top', 'single', '8'), ('bottom', 'single', '8'), ('insideH', 'single', '6')]:
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), val); el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    for tag in ('left', 'right', 'insideV'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0')
        borders.append(el)
    tblPr.append(borders)


def tri_table(doc, cap, rows, widths):
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(6)
    caption_p.paragraph_format.space_after = Pt(3)
    caption_p.paragraph_format.keep_with_next = True
    r = caption_p.add_run(cap)
    run_font(r, CN, EN, 10.5)
    tbl = doc.add_table(rows=len(rows), cols=len(widths))
    tbl.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j < 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = Pt(20)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            rr = p.add_run(val)
            run_font(rr, CN, EN, 10.5, bold=(i == 0))
    set_table_h_borders(tbl)
    return tbl


doc = Document(DOC)

doc.add_page_break()
doc.add_paragraph('第五章  系统实现', style='Heading 1')

body(doc, "本章介绍系统的具体实现。实现工作遵循前两章确定的方案，按“公共支撑—业务功能—端侧算法”三个层次展开：公共支撑部分实现登录认证、权限控制与操作日志等横切能力；业务功能部分按七个功能域实现计划制定与执行、训练与评估、内容与咨询等核心流程；端侧算法部分实现基于 MoveNet 的姿态评估流水线。开发过程借助 AI 辅助编程提升效率，并以单元测试与前后端接口联调保证实现与设计的一致性。")

# ================= 5.1 开发环境 =================
doc.add_paragraph('5.1  系统开发环境与部署', style='Heading 2')
doc.add_paragraph('5.1.1  开发环境', style='Heading 3')
body(doc, "系统采用前后端分离方式开发，后端使用 Java 语言与 Spring Boot 框架，前端使用 TypeScript 与 Vue 3 框架，端侧姿态推理基于 ONNX Runtime Web 加载 MoveNet Lightning 模型。主要开发工具与技术版本如表5-1所示。")
tri_table(doc, '表5-1  系统开发环境', [
    ('类别', '项目', '说明 / 版本'),
    ('操作系统', '开发 / 部署', 'Windows 11 开发，Linux（Docker）部署'),
    ('开发语言', '后端 / 前端', 'Java 17 / TypeScript 5'),
    ('后端框架', '核心框架', 'Spring Boot 3.2.5'),
    ('数据访问', '持久层', 'MyBatis-Plus 3.5.7 + Flyway 数据库迁移'),
    ('数据库', '关系型数据库', 'MySQL 8.0.36（InnoDB / utf8mb4）'),
    ('缓存', '内存数据库', 'Redis 7.2'),
    ('对象存储', '文件服务', 'MinIO（演示视频、内容附件）'),
    ('前端框架', '管理端 / 患者端', 'Vue 3.4.21 + Vite 5.2；Element Plus 2.7.6 / Vant 4.10.2'),
    ('姿态推理', '模型 / 运行时', 'MoveNet Lightning（192×192）ONNX / onnxruntime-web 1.30'),
    ('认证与安全', '组件', 'jjwt 0.12.6、BCrypt、EasyCaptcha 1.6.2'),
    ('接口文档', '组件', 'Knife4j 4.5.0'),
    ('工具库', '组件', 'Hutool 5.8.27'),
    ('构建与部署', '工具', 'Maven、npm + Vite、Docker Compose、Nginx'),
], (2.6, 3.4, 9.6))

body(doc, "后端采用 Maven 多模块结构，划分为公共组件、实体与数据访问、接口层等模块，其中接口层又分为患者端与管理端两个可执行模块，业务代码按七个功能域分包组织，与数据库模块一一对应；前端分为治疗师与管理端、患者端两个工程，公共的类型契约与端侧姿态评估引擎在两个工程中保持同步，避免算法口径漂移。数据库表结构由 Flyway 以版本化脚本管理，保证开发、测试与部署环境的表结构一致。")

doc.add_paragraph('5.1.2  部署与运行环境', style='Heading 3')
body(doc, "系统采用 Docker Compose 编排，在单台服务器上运行 Nginx、后端应用、MySQL、Redis 与 MinIO 五个容器。数据库、缓存与对象存储仅在容器内部网络互通，不对外暴露端口，仅由 Nginx 对外提供 80 与 443 端口，缩小了攻击面。Nginx 除分发三端静态资源、反向代理接口外，还统一注入跨源隔离响应头（COOP/COEP），为端侧 WebAssembly 推理提供所需的隔离环境。")
body(doc, "端侧姿态推理完全在用户浏览器中执行，服务器不承担图形计算压力，普通配置的云服务器即可承载三端服务。运行时，MySQL 数据目录、MinIO 对象存储目录与证书文件均挂载到宿主机数据卷，容器重建不丢失数据；系统提供定时备份与恢复机制，恢复操作需二次确认。当用户规模扩大时，可将后端应用与数据库分别横向扩展，整体架构无需调整。")

# ================= 5.2 公共功能 =================
doc.add_paragraph('5.2  系统公共功能实现', style='Heading 2')
doc.add_paragraph('5.2.1  用户登录与认证实现', style='Heading 3')
body(doc, "三端共用统一登录入口与认证服务，患者端登录界面如图5-1所示。登录流程依次执行图形验证码校验、账号锁定检查、密码校验、账号状态校验与会话签发：图形验证码由 EasyCaptcha 生成并写入 Redis，且为一次性，校验后立即删除；账号连续登录失败达到 5 次后被锁定 30 分钟，锁定状态记录于 Redis；密码采用 BCrypt 加盐散列存储，校验失败时以相同提示掩盖“账号不存在”与“密码错误”的差别，避免账号枚举；校验通过后由 jjwt 签发访问令牌，并将其唯一标识写入 Redis 会话，登出或管理员禁用账号时删除该标识即可使令牌即时失效。登录认证的核心实现如代码5-1所示。")
img_placeholder(doc, '图5-1  系统登录界面')
code_block(doc, '代码5-1  用户登录认证核心实现', [
    '// 1. 校验图形验证码（一次性，校验后立即从 Redis 删除）',
    'captchaService.verify(request.getCaptchaId(), request.getCaptchaCode());',
    '',
    '// 2. 账号锁定检查：连续失败 5 次锁定 30 分钟',
    'if (redis.hasKey(LOGIN_LOCK + username)) throw new BusinessException(ACCOUNT_LOCKED);',
    '',
    '// 3. 密码校验（账号不存在时返回相同错误，避免账号枚举）',
    'UserAccount user = userAccountMapper.selectByUsername(username);',
    'if (user == null || !passwordEncoder.matches(request.getPassword(), user.getPasswordHash())) {',
    '    handleLoginFailure(user, username, ip, deviceId);   // 累计失败次数',
    '    throw new BusinessException(LOGIN_FAILED);',
    '}',
    '',
    '// 4. 签发 JWT 并写入 Redis 会话（jti 支持登出失效与禁用强踢）',
    'LoginUser loginUser = new LoginUser(user.getUserId(), user.getUsername(), user.getRoleCode());',
    'String token = jwtTokenProvider.createAccessToken(loginUser);',
    'Claims claims = jwtTokenProvider.parse(token);',
    'redis.opsForValue().set(JWT_SESSION + user.getUserId(), claims.getId(), Duration.ofSeconds(expireSec));',
])
body(doc, "除密码登录外，系统在登录、登出、修改密码等关键动作上均记录登录日志；修改密码后删除当前会话，要求重新登录，确保凭据变更后旧令牌不再有效。")

doc.add_paragraph('5.2.2  权限控制与操作日志实现', style='Heading 3')
body(doc, "系统采用基于角色的访问控制与对象级归属校验相结合的权限模型。请求进入后端后，由拦截器解析访问令牌并构建当前登录用户对象（包含用户标识与角色），随后按接口上标注的权限注解依次执行角色校验与归属校验：患者只能访问本人的计划、任务、记录与报告，治疗师只能访问与其存在有效绑定关系的患者数据，管理员可管理账号、绑定、动作规则与公共内容。归属校验以切面方式统一实现，避免在每个业务方法中重复编写，降低越权风险。")
body(doc, "操作日志由自定义注解与事件监听机制实现：业务方法标注日志注解后，切面在方法返回时发布操作事件，由监听器异步写入操作日志表，记录操作人、操作类型、目标对象、执行结果与耗时。操作日志与登录日志共同构成完整的运维追溯链路，管理员可按操作人、类型与时间检索，用于安全审计与问题定位。")

# ================= 5.3 核心业务功能 =================
doc.add_paragraph('5.3  系统核心业务功能实现', style='Heading 2')

doc.add_paragraph('5.3.1  康复计划管理实现', style='Heading 3')
body(doc, "康复计划管理面向治疗师，是最主要的业务入口之一。治疗师登录后进入工作台，可查看其负责的患者列表与待办事项（待审核报告、待回复咨询等），界面如图5-2所示。")
img_placeholder(doc, '图5-2  康复治疗师端工作台界面')
body(doc, "制定计划时，治疗师选择患者并查阅其历史训练记录与评估报告，从标准动作库中挑选动作并设定组数、次数与难度参数，系统生成计划草稿，界面如图5-3所示。计划以“计划—计划动作项”的结构存储并带版本号：对已发布计划的修改不覆盖原计划，而是生成新版本并记录变更说明，正在执行的计划保持原版本，避免训练过程中评估标准发生变化。计划保存为草稿时不校验完整度，发布时校验动作项非空、起止日期合法、训练频率配置完整等条件，校验通过后按调度生成每日训练任务。")
img_placeholder(doc, '图5-3  康复计划制定界面')

doc.add_paragraph('5.3.2  每日训练任务与进度实现', style='Heading 3')
body(doc, "每日训练任务由康复计划按训练频率自动展开。系统按“计划动作项 + 任务日期”的复合唯一键生成任务，重复执行时幂等跳过，避免重复生成同一动作的同日任务。患者端以“今日任务”为首页核心，展示当日待训练动作、目标组数与完成状态，界面如图5-4所示。")
img_placeholder(doc, '图5-4  康复患者端每日训练任务界面')
body(doc, "患者还可通过任务日历查看整体安排，并在进度页查看训练完成情况与趋势。对确因身体原因无法完成的任务，患者可填写原因后跳过，跳过记录同步给治疗师，作为后续调整计划的依据。")

doc.add_paragraph('5.3.3  实时训练会话与记录提交实现', style='Heading 3')
body(doc, "实时训练是系统的核心场景。患者点击某个待训练任务后，后端先下发该动作的参数快照，包括评估规则、算法参数、动作签名以及动作版本所绑定的姿态引擎标识；前端随后加载本地 MoveNet 模型并对摄像头视频流进行端侧推理，界面如图5-5所示。")
img_placeholder(doc, '图5-5  实时姿态评估与训练界面')
body(doc, "训练界面实时叠加人体骨架，展示当前动作阶段、已完成次数与评分预览，并在姿态偏差超限时给出文字纠错提示。为保护隐私，原始视频与逐帧关键点数据始终保留在浏览器本地，不上传服务器。系统对姿态引擎版本做一致性校验：若端侧引擎标识与动作版本记录不一致，则中止训练并提示联系治疗师更新动作版本，确保评估标准与算法实现匹配。")
body(doc, "训练结束后，端侧汇总生成不含原始视频的评估结果并提交后端，写入训练记录与纠错日志。为避免网络抖动造成重复提交，提交接口以“每日任务标识 + 客户端请求标识”唯一键实现幂等：服务端先按该组合查询既有记录，命中则直接返回首次结果；并发提交时由数据库唯一约束兜底，捕获重复键后回查并返回首次结果。客户端在网络异常时将待提交数据暂存本地队列，网络恢复后自动重试，保证训练数据不丢失。幂等提交的核心实现如代码5-2所示。")
code_block(doc, '代码5-2  训练记录幂等提交核心实现', [
    '// 幂等：相同 (dailyTaskId, clientRequestId) 直接返回首次结果',
    'TrainingRecord existing = mapper.selectOne(new LambdaQueryWrapper<TrainingRecord>()',
    '        .eq(TrainingRecord::getDailyTaskId, request.getDailyTaskId())',
    '        .eq(TrainingRecord::getClientRequestId, request.getClientRequestId())',
    '        .last("limit 1"));',
    'if (existing != null) {',
    '    return buildSubmitVO(existing.getRecordId(), task, duplicateVO(true));',
    '}',
    '',
    'try {',
    '    trainingRecordMapper.insert(record);',
    '} catch (DuplicateKeyException e) {',
    '    // 并发同键提交：由唯一约束兜底，回查首次结果返回',
    '    TrainingRecord winner = mapper.selectOne(new LambdaQueryWrapper<TrainingRecord>()',
    '            .eq(TrainingRecord::getDailyTaskId, request.getDailyTaskId())',
    '            .eq(TrainingRecord::getClientRequestId, request.getClientRequestId())',
    '            .last("limit 1"));',
    '    if (winner != null) return buildSubmitVO(winner.getRecordId(), task, duplicateVO(true));',
    '    throw e;',
    '}',
])

doc.add_paragraph('5.3.4  训练记录与评估报告实现', style='Heading 3')
body(doc, "训练记录列表按时间倒序展示患者的历次训练，支持按日期与动作筛选；记录详情展示训练总分、六维评分、关节角度曲线以及纠错区间，界面如图5-6所示。")
img_placeholder(doc, '图5-6  训练记录详情界面')
body(doc, "评估报告以计划周期为单位自动生成。调度任务在低峰期按 ISO 周聚合各计划已完成的训练任务与训练记录，计算任务完成度、平均总分与各维度均值，固化为数据快照并预生成报告草稿；AI 依据统计快照生成小结草稿并写入草稿字段，但不绕过治疗师审核。治疗师查看报告后填写审核结论与康复建议，报告通过审核后对患者可见，界面如图5-7所示。")
img_placeholder(doc, '图5-7  评估报告详情界面')
body(doc, "若报告周期内的训练数据发生更正，治疗师可触发报告重算，系统按指定计划与周期重新聚合统计，生成版本号递增的新草稿并保留原版本，保证报告可追溯。")

doc.add_paragraph('5.3.5  动作库与评估规则维护实现', style='Heading 3')
body(doc, "标准动作与评估规则是姿态评估的基准数据，由管理员维护，采用“标准动作—动作版本—评估规则—动作模板”四层结构：标准动作记录动作名称、类型与示范视频；动作版本在创建后不可修改，版本内配置算法参数与动作签名；评估规则定义目标关节、角度区间、保持要求与评分权重；动作模板存储标准动作的关键点序列，用于相似度计算。维护界面如图5-8所示。")
img_placeholder(doc, '图5-8  动作库与评估规则维护界面')
body(doc, "为保证评分口径一致，系统在保存动作版本时校验六个维度的评分权重之和是否为 1，超出范围或权重之和不为 1 时禁止发布。动作版本一经发布即冻结，修改规则只能通过发布新版本实现，从而保证历史训练记录所参照的评估标准始终可追溯，训练开始时下发的参数快照也与动作版本一致。")

doc.add_paragraph('5.3.6  公共康复内容与在线咨询实现', style='Heading 3')
body(doc, "公共康复内容模块为患者提供康复知识中心。内容由撰稿人编辑、管理员审核后发布，关联康复部位并支持浏览、搜索与收藏，界面如图5-9所示。内容设置发布栏目与有效期，到期或被举报违规时标记下架并保留历史记录，保证内容处于受控状态。")
img_placeholder(doc, '图5-9  公共康复知识中心界面')
body(doc, "在线咨询模块提供 AI 问答与人工咨询两部分。AI 问答基于大模型接口实现，仅回答一般性康复知识并附免责声明，不输出诊断、处方与风险结论；同时按用户做每日调用限流，超出限额返回提示。对超出 AI 能力范围的问题，系统引导患者转为人工咨询，由治疗师结合患者档案回复，界面如图5-10所示。")
img_placeholder(doc, '图5-10  在线咨询与AI问答界面')

# ================= 5.4 端侧算法 =================
doc.add_paragraph('5.4  端侧姿态评估算法实现', style='Heading 2')
doc.add_paragraph('5.4.1  模型加载与端侧推理实现', style='Heading 3')
body(doc, "端侧姿态评估以 MoveNet Lightning 模型的 ONNX 文件为推理核心，由 ONNX Runtime Web 在浏览器中加载执行。模型文件与 WebAssembly 运行时均随前端静态资源本地分发，运行时不访问任何远端地址，以满足隐私保护与离线可用要求。加载时先读取模型清单（包含模型文件名、大小、SHA-256 摘要、输入输出规格与引擎标识），并对模型文件计算 SHA-256 摘要进行校验，校验不通过则中止加载并给出明确提示，避免使用被篡改的模型。")
body(doc, "推理会话按 WebGL、WebAssembly 的顺序尝试创建：优先使用 WebGL 加速，不可用时降级为 WebAssembly 后端；每次创建均设置超时上限并采用超时竞速，避免个别环境下推理线程未就绪导致整页卡死。推理前对画面按感兴趣区域外扩并做方形裁剪，缩放为模型要求的 192×192 输入；推理输出 17 个关键点的归一化坐标与置信度，再映射回原始画面，并根据是否使用前置摄像头做水平镜像，保证解剖学左右命名正确。")

doc.add_paragraph('5.4.2  关键点质量门控与时序平滑修复', style='Heading 3')
body(doc, "受光照、遮挡与快速运动影响，逐帧关键点不可避免存在噪声与缺失。系统首先对每帧做质量门控，按置信度将各关键点划分为高、中、低三级：高置信度直接采纳，中置信度采纳但计入中等质量，低置信度判为遮挡或缺失并交由修复模块处理；再依据肩、肘、髋、膝等核心关键点的可用比例判定帧质量，质量过低的帧不参与状态切换与评分。")
body(doc, "对标记为缺失的关键点，系统采用中值滤波与指数滑动平均相结合的方式平滑轨迹，并对短暂缺失（连续帧数不超过设定上限）进行保持或线性插值修复；对末端持续遮挡的关键点，则依据骨长中位数与父关节方向做几何约束预测，仅在预测值与朴素保持值偏差超过骨长容忍比例时才替换。所有插值与预测得到的点均计入补点比例，补点过多或有效帧占比过低的训练记为数据不足，不予评分，从源头抑制由遮挡引起的评分失真。")

doc.add_paragraph('5.4.3  运动特征提取与动作计数', style='Heading 3')
body(doc, "在关键点得到平滑修复后，系统进行骨架归一化并计算关节角度、关节活动度、角速度与角加速度等运动特征。动作计数由阶段状态机实现：一次标准动作被划分为就绪、抬升、保持、回落四个阶段，采用双阈值与迟滞带避免在阈值附近反复抖动，并要求状态连续保持若干帧后才确认切换；一个完整周期结束后，若幅度与保持时间达到最低门槛则计入达标次数，时长过短的动作被视为甩动而不计次。")
body(doc, "为避免误计数，系统还会将观测到的运动序列与动作签名（目标关节、起始与峰值角度区间、运动方向与阶段顺序）比对，反方向或阶段顺序不符的周期被标记为不匹配并触发提示，而不计入达标次数；若连续多个周期均不匹配，系统建议结束训练并联系治疗师，提示可能存在动作选择错误。")

doc.add_paragraph('5.4.4  六维加权评分与实时纠错', style='Heading 3')
body(doc, "训练评分从幅度、保持、节奏、稳定、对称、相似度六个维度进行。其中幅度维考察关节活动度是否达到目标；保持维考察末端姿态的保持时长；节奏维综合动作周期与期望时长的吻合度及周期间的节拍一致性；稳定维通过角度序列的高频抖动残差衡量动作平稳程度；对称维比较左右两侧同名关节的角度差异；相似度维采用动态时间规整（DTW）计算观测序列与标准模板的相似程度。各维度在周期内评分后，按时间高斯窗加权聚合为本次训练的分项得分，再按动作版本配置的权重加权求和得到总分；当模板缺失时相似度维自动降权为零，不影响整体评分。六维加权聚合的核心实现如代码5-3所示。")
code_block(doc, '代码5-3  六维加权评分聚合核心实现', [
    '// 各维度按有效周期加权平均（权重含关键帧高斯时间窗）',
    'const detail: ScoreDetail = { ...ZERO };',
    'if (!insufficient && validCycles.length > 0) {',
    '  detail.amplitude = round1(dimAvg((s) => s.amplitude) ?? 0);',
    '  detail.hold      = round1(dimAvg((s) => s.hold) ?? 0);',
    '  detail.rhythm    = round1(dimAvg((s) => s.rhythm) ?? 0);',
    '  detail.stability = round1(dimAvg((s) => s.stability) ?? 0);',
    '  detail.symmetry  = round1(dimAvg((s) => s.symmetry) ?? 0);',
    '  detail.similarity= round1(dimAvg((s) => s.similarity) ?? 0);',
    '}',
    '// 总分 = Σ(权重 × 维度分) / Σ权重（模板缺失时相似度权重自动置 0）',
    'const w = weights;',
    'const total = (detail.amplitude * w.amplitude + detail.hold * w.hold',
    '             + detail.rhythm * w.rhythm + detail.stability * w.stability',
    '             + detail.symmetry * w.symmetry + detail.similarity * w.similarity)',
    '             / effectiveWeightSum;',
    'summary.totalScore = insufficient ? 0 : round1(clamp(total, 0, 100));',
])
body(doc, "上述门控、平滑与评分所涉及的算法参数集中记录于动作版本中，端侧引擎在训练开始时随参数快照下发，其默认取值如表5-2所示。")
tri_table(doc, '表5-2  端侧姿态评估关键算法参数', [
    ('参数', '默认取值', '说明'),
    ('confHigh', '0.5', '关键点高置信度阈值，高于该值直接采纳'),
    ('confLow', '0.2', '关键点低置信度阈值，低于该值判为遮挡 / 缺失'),
    ('medianWindow', '3', '中值滤波窗口（帧，奇数）'),
    ('emaAlpha', '0.35', '指数滑动平均系数，越大越跟手'),
    ('maxGapFrames', '5', '允许连续补点的最大帧数，超过则判该段无效'),
    ('boneLenTolerance', '0.15', '骨长一致性容忍比例，用于几何修复判定'),
    ('repairRatioLimit', '0.3', '单段训练允许的补点占比上限'),
    ('stateHoldFrames', '3', '状态切换需连续确认的帧数（去抖）'),
    ('hysteresisBand', '8', '状态切换迟滞带（度）'),
    ('minCycleSec', '1.2', '最小有效动作周期时长（秒），短于该值判为甩动'),
    ('validFrameRatioMin', '0.7', '有效帧占比下限，低于该值判为数据不足'),
    ('六维权重', '0.30 / 0.20 / 0.15 / 0.15 / 0.10 / 0.10', '幅度 / 保持 / 节奏 / 稳定 / 对称 / 相似度，合计为 1'),
], (3.6, 4.8, 7.2))
body(doc, "实时纠错根据当前帧的关节角度偏差、抖动与左右对称差异生成提示：偏差在软阈值内不提示，超出软阈值提示一级（黄色），超出硬阈值提示二级（红色）。提示按“严重度 × 规则权重 × 归一化偏差”排序，同屏最多显示 1 至 2 条，同一规则相邻区间三秒内不重复弹出，持续偏差区间则常驻显示。训练结束后，系统按“规则 × 连续区间”将提示聚合为纠错区间，随训练记录一并保存，供治疗师分析患者的常见问题并据此调整康复计划。")

doc.save(DOC)
print('saved:', DOC)
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
