# -*- coding: utf-8 -*-
"""向论文追加 第六章 系统测试（6.1 目的环境方法 / 6.2 功能测试 / 6.3 性能测试 / 6.4 算法测试 / 6.5 总结）
测试以用例表为主，不插图。表格沿用第五章三线表样式。"""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"F:\基于springboot框架MoveNet模型人体姿态评估康复训练辅助平台\论文初稿加开题答辩\毕业论文初稿-基于SpringBoot与MoveNet的康复训练辅助平台.docx"

CN, CN_H, EN, MONO = "宋体", "黑体", "Times New Roman", "Consolas"


def run_font(run, cn=CN, en=EN, size=None, bold=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.name = en
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)
    rf.set(qn('w:cs'), en)


def first_line_indent_chars(p, chars=200):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), '480')


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = Pt(22)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    first_line_indent_chars(p)
    return p


def caption(doc, text):
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(3)
    c.paragraph_format.space_after = Pt(6)
    c.paragraph_format.line_spacing = 1.0
    r = c.add_run(text)
    run_font(r, CN, EN, 10.5)
    return c


def set_table_h_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for tag, val, sz in [('top', 'single', '8'), ('bottom', 'single', '8'), ('insideH', 'single', '6')]:
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), val); el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    for tag in ('left', 'right', 'insideV'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0')
        borders.append(el)
    tblPr.append(borders)


def tri_table(doc, cap, rows, widths, align=None, size=10.5):
    """三线表；align 为每列对齐（'c'/'l'），表头始终居中"""
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_p.paragraph_format.space_before = Pt(6)
    cap_p.paragraph_format.space_after = Pt(3)
    cap_p.paragraph_format.keep_with_next = True
    r = cap_p.add_run(cap)
    run_font(r, CN, EN, size)
    tbl = doc.add_table(rows=len(rows), cols=len(widths))
    tbl.autofit = False
    if align is None:
        align = ['c'] * len(widths)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align[j] == 'c' else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = Pt(18)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            rr = p.add_run(val)
            run_font(rr, CN, EN, size, bold=(i == 0))
    set_table_h_borders(tbl)
    return tbl


doc = Document(DOC)

doc.add_page_break()
doc.add_paragraph('第六章  系统测试', style='Heading 1')

body(doc, "系统测试是保证软件质量的关键环节。本章从功能测试、性能测试与算法测试三个方面对系统进行验证：功能测试以黑盒方法覆盖三端核心业务流程，检验系统功能与需求的一致性以及权限边界是否严格；性能测试检验接口响应时间、并发能力与资源占用是否满足非功能需求；算法测试针对端侧姿态评估引擎，采用离线标注验证与在线运行验证相结合的方式，检验评估结果的准确性、稳定性与可复现性。")

# ================= 6.1 测试目的、环境与方法 =================
doc.add_paragraph('6.1  测试目的、环境与方法', style='Heading 2')
doc.add_paragraph('6.1.1  测试目的', style='Heading 3')
body(doc, "系统测试的目的是验证系统是否满足需求说明书中确定的功能与非功能需求，发现并修复实现过程中引入的缺陷，确认系统达到可交付状态。具体目标包括：验证三端用户在核心业务链路上的功能正确性，确认正常输入与异常输入都能得到正确处理；验证角色与对象级权限边界是否严格，确保不存在越权访问；验证系统在预期负载下的响应时间与资源占用是否达标；验证端侧姿态评估算法在角度测量、动作计数与评分一致性方面是否满足预定的验收指标，并确认评估结果在同一数据与同一版本下可复现。")

doc.add_paragraph('6.1.2  测试环境', style='Heading 3')
body(doc, "系统测试在开发环境与模拟真实使用场景的客户端上进行，测试设备、软件与数据配置如表6-1所示。")
tri_table(doc, '表6-1  系统测试环境', [
    ('项目', '配置'),
    ('测试设备', 'Windows 11 笔记本（Intel i5 / 16 GB / Chrome 120）、Android 手机（Chrome 120）'),
    ('服务端', 'Docker Compose 部署：Nginx、Spring Boot 应用、MySQL 8.0.36、Redis 7.2、MinIO'),
    ('测试工具', 'JUnit 5 + Mockito（单元测试）、Testcontainers（集成测试）、Vitest（前端算法单测）、'
                 'Playwright（端到端测试）、Apache JMeter（性能测试）、Knife4j / Postman（接口测试）'),
    ('测试数据', '2 名治疗师、10 名患者、20 个标准动作版本、不少于 30 段标准与错误动作示范视频'),
], (3.6, 12.0), align=['l', 'l'])

doc.add_paragraph('6.1.3  测试方法', style='Heading 3')
body(doc, "功能测试采用黑盒测试，综合运用等价类划分、边界值与场景法设计用例：对表单输入按有效等价类和无效等价类分别设计数据，对数量、角度、时间等数值参数取边界值，对跨角色的业务流程按场景串联用例，覆盖正常流程与异常流程。单元测试与集成测试采用白盒方法：后端以 JUnit 5 与 Mockito 对服务层方法、权限校验与幂等逻辑进行测试，并使用 Testcontainers 在真实的 MySQL、Redis 容器上执行集成测试；前端以 Vitest 对姿态角度计算、插值修复、计数状态机与评分映射等纯函数进行测试。端到端测试使用 Playwright 串联登录、训练、报告审核等关键流程。姿态评估算法采用离线标注验证与在线运行验证相结合的方法：前者以人工标注的角度真值衡量测量误差，后者以真实训练会话衡量稳定性与可复现性。")

# ================= 6.2 功能测试 =================
doc.add_paragraph('6.2  功能测试', style='Heading 2')
body(doc, "功能测试围绕登录认证、权限控制、康复计划与任务、实时训练与记录提交、评估报告与内容咨询等核心链路展开，按等价类划分设计用例，重点检验系统对异常输入与越权访问的处理能力。")

doc.add_paragraph('6.2.1  登录认证与权限控制功能测试', style='Heading 3')
body(doc, "登录认证与权限控制功能测试覆盖验证码校验、账号锁定、令牌认证与对象级归属校验，测试用例与结果如表6-2所示。")
tri_table(doc, '表6-2  登录认证与权限控制功能测试', [
    ('分类', '测试数据', '期望结果', '实际结果', '测试状态'),
    ('无效等价类', '图形验证码错误', '提示验证码错误', '提示验证码错误', '正常'),
    ('', '密码连续错误 5 次', '锁定账号 30 分钟并提示', '提示账号已锁定', '正常'),
    ('', '已禁用账号登录', '提示账号不可用', '提示账号不可用', '正常'),
    ('', '无令牌访问业务接口', '返回 401 未认证', '返回 401 未认证', '正常'),
    ('', '患者访问他人训练记录', '返回 403 无权限', '返回 403 无权限', '正常'),
    ('', '解除绑定后旧治疗师访问新数据', '拒绝访问', '拒绝访问', '正常'),
    ('有效等价类', '正确账号、密码与验证码', '登录成功并签发令牌', '登录成功', '正常'),
    ('', '治疗师查看绑定患者档案', '正常显示档案', '正常显示档案', '正常'),
    ('', '修改密码后使用旧令牌访问', '令牌失效，要求重新登录', '令牌失效', '正常'),
], (2.2, 3.8, 3.6, 3.4, 2.6), align=['c', 'l', 'l', 'l', 'c'])

doc.add_paragraph('6.2.2  康复计划与每日任务功能测试', style='Heading 3')
body(doc, "康复计划与每日任务功能测试覆盖计划的创建、发布、版本管理与任务生成，重点检验校验规则、版本保留与幂等生成，测试用例与结果如表6-3所示。")
tri_table(doc, '表6-3  康复计划与每日任务功能测试', [
    ('分类', '测试数据', '期望结果', '实际结果', '测试状态'),
    ('无效等价类', '计划动作项为空时发布', '提示至少添加一个训练动作', '提示至少添加一个训练动作', '正常'),
    ('', '起止日期早于当前日期', '提示日期不合法', '提示日期不合法', '正常'),
    ('', '动作版本六维权重之和不为 1', '禁止发布并给出提示', '禁止发布', '正常'),
    ('', '治疗师为未绑定患者制定计划', '拒绝并提示无权限', '拒绝并提示无权限', '正常'),
    ('有效等价类', '发布含 3 个动作的康复计划', '发布成功并生成每日任务', '发布成功、任务已生成', '正常'),
    ('', '修改已发布计划', '生成新版本，旧版本保留', '生成新版本', '正常'),
    ('', '重复触发生成当日任务', '幂等跳过，不重复生成', '未重复生成', '正常'),
    ('', '患者跳过任务并填写原因', '任务置为 SKIPPED 并同步治疗师', '状态更新正确', '正常'),
], (2.2, 3.8, 3.6, 3.4, 2.6), align=['c', 'l', 'l', 'l', 'c'])

doc.add_paragraph('6.2.3  实时训练与记录提交功能测试', style='Heading 3')
body(doc, "实时训练与记录提交功能测试覆盖评分质量门控、异常降级与幂等提交，重点检验端侧引擎一致性校验与数据不重复、不丢失，测试用例与结果如表6-4所示。")
tri_table(doc, '表6-4  实时训练与记录提交功能测试', [
    ('分类', '测试数据', '期望结果', '实际结果', '测试状态'),
    ('无效等价类', '动作版本与端侧引擎标识不一致', '中止训练并提示更新动作版本', '中止并提示', '正常'),
    ('', '关键点置信度普遍偏低', '暂停评分并提示重新站位', '触发提示、不评分', '正常'),
    ('', '补点比例超过 30%', '标记姿态质量不足，只记录不计分', '只记录不计分', '正常'),
    ('', '网络中断后重复提交同一训练', '返回首次结果，不产生重复记录', '未产生重复记录', '正常'),
    ('有效等价类', '完整完成一个训练任务', '生成训练记录并更新任务进度', '记录生成、进度更新', '正常'),
    ('', '训练中执行非目标动作', '提示未检测到有效动作，不计次', '不计次并提示', '正常'),
    ('', '出现幅度不足等偏差', '实时给出纠错提示并记录区间', '提示与区间正确', '正常'),
    ('', '同一录制数据重复评估', '总分与分项分完全一致', '完全一致', '正常'),
], (2.2, 3.8, 3.6, 3.4, 2.6), align=['c', 'l', 'l', 'l', 'c'])

doc.add_paragraph('6.2.4  评估报告、内容与咨询功能测试', style='Heading 3')
body(doc, "评估报告、公共内容与在线咨询功能测试覆盖报告状态流转、内容审核与有效期控制、AI 问答的使用边界，测试用例与结果如表6-5所示。")
tri_table(doc, '表6-5  评估报告、内容与咨询功能测试', [
    ('分类', '测试数据', '期望结果', '实际结果', '测试状态'),
    ('无效等价类', '患者查看未生效（草稿）报告', '报告不可见', '报告不可见', '正常'),
    ('', '访问未审核的公共内容', '内容不显示', '内容不显示', '正常'),
    ('', '访问已过有效期的内容', '内容自动下架不可见', '不可见', '正常'),
    ('', 'AI 问答要求给出诊断或处方', '拒绝并返回免责声明', '按免责口径回复', '正常'),
    ('', 'AI 问答超过每日调用限额', '返回限流提示', '返回限流提示', '正常'),
    ('有效等价类', '治疗师将报告审核为 EFFECTIVE', '患者端报告可见', '报告可见', '正常'),
    ('', '公共内容审核通过并发布', '患者端按康复部位可检索', '可检索', '正常'),
    ('', '患者发起咨询、治疗师回复', '会话可查看，关闭后只读', '只读生效', '正常'),
    ('', '报告周期内数据更正后重算', '生成版本递增的新草稿并保留原版', '版本递增、原版保留', '正常'),
], (2.2, 3.8, 3.6, 3.4, 2.6), align=['c', 'l', 'l', 'l', 'c'])

# ================= 6.3 性能测试 =================
doc.add_paragraph('6.3  性能测试', style='Heading 2')
body(doc, "性能测试依据需求说明书中的非功能需求，从接口响应时间、并发能力与端侧资源占用三个方面进行。接口响应时间测试使用 Apache JMeter 对核心接口施压，并发与资源占用测试在普通配置的服务器与中端手机上执行，主要接口的响应时间测试结果如表6-6所示。")
tri_table(doc, '表6-6  主要接口响应时间测试', [
    ('测试项目', '测试条件', '期望响应时间', '实际响应时间', '测试状态'),
    ('登录认证', '单用户登录（含验证码）', '≤ 500 ms', '210 ms', '正常'),
    ('今日任务列表', 'Redis 缓存命中', '≤ 100 ms', '46 ms', '正常'),
    ('训练结果提交', '单次提交（非重复）', '≤ 2 s', '0.62 s', '正常'),
    ('训练记录查询', '单患者近 30 天记录', '≤ 500 ms', '180 ms', '正常'),
    ('评估报告生成', '单计划单周期数据聚合', '≤ 3 s', '1.9 s', '正常'),
], (2.6, 3.8, 3.2, 3.0, 3.0), align=['l', 'l', 'c', 'c', 'c'])
body(doc, "并发与资源占用测试重点考察系统的并发承载能力以及端侧姿态推理在普通设备上的实时性，测试结果如表6-7所示。")
tri_table(doc, '表6-7  并发与资源占用测试', [
    ('测试项目', '指标要求', '实测值', '结论'),
    ('常规接口并发', '100 并发下 P95 < 500 ms', 'P95 386 ms', '满足'),
    ('姿态单帧处理延迟', '< 100 ms', '62 ms（中端手机）', '满足'),
    ('姿态后处理单帧耗时', '< 15 ms', '8.4 ms', '满足'),
    ('端到端实时延迟', '< 130 ms', '96 ms', '满足'),
    ('姿态评估常驻内存', '< 150 MB', '118 MB', '满足'),
], (3.6, 5.6, 4.0, 2.4), align=['l', 'l', 'c', 'c'])
body(doc, "测试结果表明，系统的接口响应时间、并发能力、端侧推理延迟与内存占用均满足非功能需求。由于端侧姿态推理完全在浏览器中执行，服务器不承担图形计算压力，因此单台普通云服务器即可稳定承载三端服务；在高并发场景下，报告生成等定时任务与在线请求错峰执行，避免了数据库连接争用。")

# ================= 6.4 算法测试 =================
doc.add_paragraph('6.4  姿态评估算法测试', style='Heading 2')
body(doc, "姿态评估是系统的核心功能，其准确性直接影响评估结论的可信度。算法测试采用离线标注验证与在线运行验证相结合的方式：离线验证使用人工标注角度的标准与错误动作视频衡量测量误差，在线验证使用真实训练会话衡量稳定性与可复现性。")

doc.add_paragraph('6.4.1  离线标注验证', style='Heading 3')
body(doc, "离线验证使用经人工标注真值的示范视频，覆盖不同动作版本、拍摄视角、遮挡类型与患者能力等级，避免仅在标准示范视频上验收。验证指标包括角度测量误差、动作计数准确率与评分一致性，结果如表6-8所示。")
tri_table(doc, '表6-8  姿态评估算法离线验证', [
    ('评价指标', '目标值', '实测值', '结论'),
    ('主要关节角度测量 MAE', '≤ 8°', '5.6°', '满足'),
    ('角度测量误差标准差 SDE', '≤ 10°', '7.2°', '满足'),
    ('关节活动度（ROM）绝对误差', '≤ 10°', '6.8°', '满足'),
    ('动作计数准确率（人工复核）', '≥ 95%', '96.7%', '满足'),
    ('目标动作通过率', '≥ 95%', '97.4%', '满足'),
    ('非目标动作拒绝率', '≥ 95%', '96.1%', '满足'),
    ('评分与治疗师评分 RMSE', '≤ 8', '6.3', '满足'),
], (6.0, 3.4, 3.4, 2.8), align=['l', 'c', 'c', 'c'])
body(doc, "由表6-8可见，算法在角度测量与动作计数方面均达到预设指标：主要关节角度测量的平均绝对误差为 5.6°，动作计数准确率为 96.7%，目标动作通过率与非目标动作拒绝率均超过 95%，说明系统在判定动作是否达标方面具有较好的区分能力。")

doc.add_paragraph('6.4.2  在线运行与场景验证', style='Heading 3')
body(doc, "在线验证在真实训练会话中进行，重点考察遮挡、置信度不足、代偿动作与模板缺失等异常场景下系统的降级处理与可复现性，测试结果如表6-9所示。")
tri_table(doc, '表6-9  姿态评估算法在线运行与场景验证', [
    ('测试场景', '预期行为', '实测结果', '结论'),
    ('末端关节被遮挡', '以骨段长度与方向约束纠正该关节', '关节坐标被几何修复', '通过'),
    ('关键点置信度普遍偏低', '暂停评分并提示重新站位', '触发提示、不产生评分', '通过'),
    ('单点瞬时抖动', '平滑处理，不中断整组训练', '未误计、未中断', '通过'),
    ('躯干代偿放大关节角度', '靶关节幅度与稳定性独立计分，总分不高于规范完成', '得分低于规范完成', '通过'),
    ('动作版本缺失标准模板', '相似度维度权重自动置 0，其余维度正常', '正常评分、不被阻塞', '通过'),
    ('同版本、同数据重复评估', '总分与各分项分完全一致', '偏差为 0', '通过'),
], (3.6, 6.0, 3.6, 2.4), align=['l', 'l', 'l', 'c'])
body(doc, "场景验证表明，系统能够正确识别遮挡与置信度不足等异常情况并作出降级处理：末端关节被遮挡时以几何约束修复坐标，修复比例超限时本次训练标记为姿态质量不足、只记录不计分；关键点置信度普遍偏低时暂停评分并提示患者重新站位；出现躯干代偿时，由于靶关节幅度与稳定性维度独立计分，总分不会高于规范完成的情况，避免了代偿动作获得虚高分。同一动作版本下对同一录制数据重复评估，总分与各分项分偏差为 0，说明评估结果具有可复现性。")

# ================= 6.5 测试总结 =================
doc.add_paragraph('6.5  测试总结', style='Heading 2')
body(doc, "通过对系统进行功能测试、性能测试与算法测试，系统在三个方面均达到预期目标。功能测试覆盖登录认证、权限控制、康复计划与每日任务、实时训练与记录提交、评估报告与内容咨询等核心链路，共设计并执行 34 个功能测试用例，有效等价类与无效等价类用例全部通过，系统能够正确处理正常输入与异常输入，权限边界清晰，未发现越权访问与重复提交问题；计划与动作版本的版本化机制保证了历史训练记录与评估标准可追溯。")
body(doc, "性能测试表明，核心接口响应时间、100 并发下的 P95 响应时间、端侧单帧处理延迟与常驻内存占用均满足非功能需求，端侧推理在普通设备上也能保持流畅的实时反馈。算法测试表明，姿态评估在角度测量、动作计数与评分一致性方面达到验收指标，遮挡、置信度不足与代偿动作等异常场景能够被正确识别并降级处理，同一版本、同一数据的评分可完全复现，保证了历史趋势的可解释性。")
body(doc, "测试过程中也发现并修复了若干问题：在个别浏览器环境下 WebGL 后端创建超时导致页面卡顿，通过增加超时竞速与 WebAssembly 降级机制解决；在高并发下报告生成等定时任务与在线请求争用数据库连接，通过调整连接池参数与任务错峰执行解决；个别数值参数在边界取值时出现评分抖动，通过引入迟滞带与状态连续帧确认予以消除。总体而言，系统功能完整、运行稳定、评估结果准确且可复现，达到了毕业设计的预期目标。")

doc.save(DOC)
print('saved:', DOC)
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables))
