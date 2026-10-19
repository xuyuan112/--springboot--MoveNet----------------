# -*- coding: utf-8 -*-
"""向论文追加 第四章 系统设计（4.1 总体架构 / 4.2 功能模块 / 4.3 业务流程 / 4.4 数据库设计）"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"F:\基于springboot框架MoveNet模型人体姿态评估康复训练辅助平台\论文初稿加开题答辩\毕业论文初稿-基于SpringBoot与MoveNet的康复训练辅助平台.docx"
IMG = r"F:\基于springboot框架MoveNet模型人体姿态评估康复训练辅助平台\论文初稿加开题答辩\毕业论文插图"

CN, CN_H, EN = "宋体", "黑体", "Times New Roman"


def run_font(run, cn=CN, en=EN, size=None, bold=None):
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    run.font.name = en
    rpr = run._r.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.insert(0, rf)
    for a in list(rf.attrib):
        del rf.attrib[a]
    rf.set(qn('w:ascii'), en)
    rf.set(qn('w:hAnsi'), en)
    rf.set(qn('w:eastAsia'), cn)
    rf.set(qn('w:cs'), en)


def first_line_indent_chars(p, chars=200):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn('w:ind'))
    if ind is None:
        ind = OxmlElement('w:ind')
        pPr.append(ind)
    ind.set(qn('w:firstLineChars'), str(chars))
    ind.set(qn('w:firstLine'), '480')


def body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    first_line_indent_chars(p)
    return p


def figure(doc, path, cap, width_cm=10.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0  # 单倍行距，避免固定值22磅裁剪内嵌图片
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_before = Pt(3)
    c.paragraph_format.space_after = Pt(6)
    r = c.add_run(cap)
    run_font(r, CN, EN, 10.5)
    return c


def set_table_h_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for tag, val, sz in [('top', 'single', '8'), ('bottom', 'single', '8'), ('insideH', 'single', '6')]:
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), val); el.set(qn('w:sz'), sz)
        el.set(qn('w:space'), '0'); el.set(qn('w:color'), '000000')
        borders.append(el)
    for tag in ('left', 'right', 'insideV'):
        el = OxmlElement('w:' + tag)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0')
        borders.append(el)
    tblPr.append(borders)


def db_table(doc, cap, rows, widths=(4.2, 3.6, 7.8)):
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_before = Pt(6)
    caption_p.paragraph_format.space_after = Pt(3)
    caption_p.paragraph_format.keep_with_next = True
    r = caption_p.add_run(cap)
    run_font(r, CN, EN, 10.5)
    tbl = doc.add_table(rows=len(rows), cols=3)
    tbl.autofit = False
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (i == 0 or j < 2) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = Pt(20)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            rr = p.add_run(val)
            run_font(rr, CN, EN, 10.5, bold=(i == 0))
    set_table_h_borders(tbl)
    return tbl


doc = Document(DOC)

doc.add_page_break()
doc.add_paragraph('第四章  系统设计', style='Heading 1')

# ---------------- 4.1 总体架构设计 ----------------
doc.add_paragraph('4.1  系统总体架构设计', style='Heading 2')
doc.add_paragraph('4.1.1  系统总体架构', style='Heading 3')
for t in [
    "平台采用前后端分离的B/S架构，自上而下划分为浏览器层、网关层、应用层、访问层与数据层五个层次，系统总体架构如图4-1所示。与一般管理类系统不同，本平台将人体姿态推理下沉到浏览器端执行，因此在浏览器层中除三端页面外，还包含一个端侧姿态推理模块，形成“端侧推理＋服务端业务支撑”的整体结构。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-1 系统总体架构图.png'), '图4-1  系统总体架构图', 14.5)
for t in [
    "浏览器层面向三类角色提供独立的页面：康复患者端基于Vant组件库构建，适配手机屏幕，承载计划查看、实时训练与报告查看等高频功能；康复治疗师端与系统管理端基于Element Plus组件库构建，以列表、表单与图表为主。端侧姿态推理部分在浏览器本地完成摄像头视频流的逐帧处理，依次执行MoveNet Lightning模型推理、关键点质量门控、时序平滑与遮挡修复、骨架归一化与运动特征计算，以及动作计数、六维加权评分与纠错提示，原始视频与逐帧关键点数据不出端，仅训练结束后向服务端上报汇总指标。",
    "网关层由Nginx承担，负责三端静态资源的分发、向应用层的反向代理、HTTPS加密传输，并统一注入COOP/COEP跨源隔离响应头，以满足ONNX Runtime Web启用多线程WebAssembly后端所需的跨源隔离环境。",
    "应用层基于Spring Boot 3框架运行于内嵌Tomcat容器，由安全与公共支撑、七大业务域两部分组成。安全与公共支撑以横切方式服务于各业务域，包括JWT身份认证、基于角色与对象级归属的权限控制、BCrypt密码散列、Redis缓存与限流、定时任务调度以及操作日志审计；七大业务域即第三章所述的账户与权限、动作与算法、康复计划、姿态评估与训练、数据与报告、公共内容与文件、咨询与运维，各业务域内部按Controller、Service、Mapper三层组织代码。",
    "访问层对下屏蔽数据存储细节：MyBatis-Plus负责MySQL的关系数据访问并提供通用增删改查与分页能力，Redis客户端负责缓存与登录限流数据的读写，MinIO SDK负责演示视频、内容附件等文件的存取。数据层由MySQL 8.0、Redis 7.2与MinIO三部分组成，分别承担业务数据存储、热点数据缓存与对象文件存储职责。",
]:
    body(doc, t)

doc.add_paragraph('4.1.2  系统部署架构', style='Heading 3')
for t in [
    "系统部署采用Docker Compose统一编排，在单台云服务器上运行Nginx、后端服务、MySQL、Redis与MinIO五个容器。容器之间通过内部网络通信，数据库、缓存与对象存储不直接对外暴露端口，仅由Nginx对外提供80与443端口，缩小了攻击面。",
    "数据持久化方面，MySQL数据目录、MinIO对象存储目录与Nginx证书均挂载到宿主机数据卷，容器重建不丢失数据；系统提供定时备份脚本，定期将业务数据与对象存储快照备份到独立存储位置，恢复操作需二次确认。姿态推理在用户浏览器端完成，服务器无图形计算压力，普通配置的云服务器即可承载，后续用户规模扩大时可将后端服务与数据库分别迁移或横向扩展，架构上无需调整。",
]:
    body(doc, t)

# ---------------- 4.2 功能模块设计 ----------------
doc.add_paragraph('4.2  系统功能模块设计', style='Heading 2')
for t in [
    "在第三章功能分析的基础上，系统自顶向下按照业务职责划分为七个功能域，每个功能域内部再按具体功能点组织，形成“系统—功能域—功能点”的三级结构，如图4-2所示。模块划分遵循高内聚、低耦合的原则，七个功能域与数据库的七个模块一一对应，域之间仅通过明确的接口协作，便于分模块开发与测试。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-2 系统功能模块图.png'), '图4-2  系统功能模块图', 15.0)
for t in [
    "从使用者的角度看，七大功能域的功能按角色组织为康复患者端、康复治疗师端与系统管理端三个入口，三端功能结构如图4-3所示。患者端聚焦计划执行与训练反馈，治疗师端聚焦方案制定与结果审核，管理端承担账号、规则、内容与运维等支撑工作；同一功能域的功能可能分布在多个端，例如数据与报告域的训练记录查看在患者端，而报告审核在治疗师端。两种划分视角相互对应，前者保证业务数据的一致与完整，后者保证各角色操作的边界清晰。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-3 系统功能模块图（按角色三端）.png'), '图4-3  系统功能模块图（按角色）', 15.0)

# ---------------- 4.3 业务流程设计 ----------------
doc.add_paragraph('4.3  系统业务流程设计', style='Heading 3')
doc.add_paragraph('4.3.1  康复计划制定流程', style='Heading 3')
for t in [
    "康复计划制定是训练执行的起点，流程如图4-4所示。治疗师在工作台中选择负责的患者，查阅其历史训练记录与评估报告，确定本次训练部位与康复目标；随后从标准动作库中挑选动作并设定组数、次数与难度参数，系统生成计划草稿；治疗师确认发布前可反复调整动作与训练参数，确认后系统创建新的计划版本并记录版本号与变更说明，计划自生效日期起对患者可见。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-4 康复计划制定流程图.png'), '图4-4  康复计划制定流程图', 8.5)

doc.add_paragraph('4.3.2  实时姿态评估与训练流程', style='Heading 3')
for t in [
    "实时姿态评估与训练是系统的核心流程，完全在浏览器端完成，流程如图4-5所示。患者授权摄像头并进入当日训练任务后，端侧MoveNet模型逐帧提取人体17个关键点；系统判断关键点置信度是否达标，不达标帧经过时序平滑与遮挡修复后仅作记录、不参与评分；达标帧进行骨架归一化并计算运动特征，再匹配当前动作版本的评估规则计算六个维度的偏差，经加权评分后实时给出纠错提示；上述过程循环执行直至本次训练结束，最后生成不含原始视频的训练汇总数据提交后端，写入训练记录与纠错日志。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-5 实时姿态评估与训练流程图.png'), '图4-5  实时姿态评估与训练流程图', 9.5)

doc.add_paragraph('4.3.3  评估报告生成与审核流程', style='Heading 3')
for t in [
    "评估报告按计划周期自动生成并经人工审核后生效，流程如图4-6所示。系统在低峰期按计划周期聚合训练记录与纠错日志，计算任务完成度、训练总分与各维度均值，生成数据快照并预生成报告草稿；治疗师查看报告与训练明细，填写审核结论与康复建议；审核不通过时补充训练数据重新聚合，通过后报告置为已生效并对患者可见；系统同时判断是否需要调整康复计划，需要时创建新的计划版本，进入下一轮训练循环。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-6 评估报告生成与审核流程图.png'), '图4-6  评估报告生成与审核流程图', 8.5)

doc.add_paragraph('4.3.4  动作库与评估规则维护流程', style='Heading 3')
for t in [
    "动作库与评估规则是姿态评估的基准数据，其维护流程如图4-7所示。管理员录入标准动作基本信息并将示范视频上传至对象存储，随后为该动作创建不可修改的动作版本；在版本内配置关键点定义、角度阈值与评分权重，系统校验同一版本内六个维度的规则权重之和是否为1，不通过则调整权重取值，通过后动作版本生效，供康复计划引用。规则调整只能通过发布新版本实现，保证历史训练数据所参照的评估标准可追溯。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-7 动作库与评估规则维护流程图.png'), '图4-7  动作库与评估规则维护流程图', 8.5)

doc.add_paragraph('4.3.5  在线咨询与AI问答流程', style='Heading 3')
for t in [
    "在线咨询与AI问答流程如图4-8所示。患者在存在有效绑定关系的前提下发起咨询，系统识别问题并检索公共康复知识，若AI能够独立解答则生成AI回复并记录会话消息，患者可继续追问；若超出AI能力范围，系统创建人工咨询工单并通知治疗师，由治疗师结合患者档案进行回复。会话关闭后双方转为只读，作为康复过程的沟通记录留存。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-8 在线咨询与AI问答流程图.png'), '图4-8  在线咨询与AI问答流程图', 9.0)

doc.add_paragraph('4.3.6  公共内容审核发布流程', style='Heading 3')
for t in [
    "公共康复内容的审核发布流程如图4-9所示。撰稿人编辑内容并关联康复部位后提交审核，管理员对内容质量与风险提示进行审核，不通过则退回并附修改意见；通过后设置发布栏目与生效期限，内容发布至公共康复知识中心供患者浏览、搜索与收藏；内容到期或发现违规时被标记下架并保留历史记录，保证内容始终处于受控状态。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-9 公共内容审核发布流程图.png'), '图4-9  公共内容审核发布流程图', 8.5)

# ---------------- 4.4 数据库设计 ----------------
doc.add_paragraph('4.4  数据库设计', style='Heading 2')
doc.add_paragraph('4.4.1  概念结构设计', style='Heading 3')
for t in [
    "根据需求分析中各业务域的实体与联系，系统抽象出用户账号、患者档案、治疗师档案、患者治疗师绑定、标准动作、动作版本、评估规则、康复计划、计划动作项、每日训练任务、训练记录、评估报告、公共内容、咨询会话等核心实体。系统总体E-R图只刻画核心实体及其主要联系，如图4-10所示：管理员建立患者与治疗师的绑定关系，治疗师基于绑定关系为患者制定康复计划，计划包含若干动作项并按日展开为训练任务，训练产生训练记录，记录按周期聚合为评估报告，形成完整的业务主链路。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-10 系统总体E-R图.png'), '图4-10  系统总体E-R图', 14.5)
for t in [
    "总体E-R图之外，再按七个业务域分别给出局部E-R图，以表达各域内实体的完整属性与联系。",
    "账户与权限域的局部E-R图如图4-11所示。用户账号作为所有关系的根实体，通过角色区分患者、治疗师与管理员；患者档案、治疗师档案分别从属于用户账号，患者与治疗师之间通过绑定实体建立多对多联系，并记录关系类型与生效区间，患者康复部位记录关联康复部位字典。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-11 账户与权限域E-R图.png'), '图4-11  账户与权限域局部E-R图', 15.0)
for t in [
    "动作与算法域的局部E-R图如图4-12所示。标准动作与动作版本为一对多联系，动作版本一经发布不可修改；每个动作版本拥有各自的评估规则与动作模板，模板下再挂接模板采样数据，共同构成姿态评估的基准。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-12 动作与算法域E-R图.png'), '图4-12  动作与算法域局部E-R图', 15.0)
for t in [
    "康复计划域的局部E-R图如图4-13所示。康复计划由计划动作项有序组成，计划调度配置训练频率，系统按调度将计划动作项展开为按日的训练任务，任务与计划动作项之间通过复合唯一键保证同一动作同一天只生成一条任务。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-13 康复计划域E-R图.png'), '图4-13  康复计划域局部E-R图', 15.0)
for t in [
    "姿态评估与训练域的局部E-R图如图4-14所示。每日训练任务与训练记录为一对多联系，训练记录固化提交时的动作版本快照与评分明细，其下再关联逐次动作的纠错日志，用于还原训练过程中的偏差细节。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-14 姿态评估与训练域E-R图.png'), '图4-14  姿态评估与训练域局部E-R图', 13.5)
for t in [
    "数据与报告域的局部E-R图如图4-15所示。评估报告从属于康复计划并按报告周期唯一，报告内嵌训练数据快照与分项统计，由治疗师审核后生效；报告的聚合来源为其周期内的训练记录。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-15 数据与报告域E-R图.png'), '图4-15  数据与报告域局部E-R图', 15.0)
for t in [
    "公共内容与文件域的局部E-R图如图4-16所示。公共内容由作者撰写、管理员审核，关联康复部位字典并支持患者收藏；内容附件与动作演示素材统一登记为文件对象，通过预签名地址访问。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-16 公共内容与文件域E-R图.png'), '图4-16  公共内容与文件域局部E-R图', 15.0)
for t in [
    "咨询与运维域的局部E-R图如图4-17所示。咨询会话基于有效的患者治疗师绑定建立，会话下挂多条会话消息与咨询附件；公告面向指定范围发布，操作日志记录关键业务的操作人与操作内容，支撑安全审计。",
]:
    body(doc, t)
figure(doc, os.path.join(IMG, '图4-17 咨询与运维域E-R图.png'), '图4-17  咨询与运维域局部E-R图', 15.0)

doc.add_paragraph('4.4.2  逻辑结构设计', style='Heading 3')
for t in [
    "将E-R模型转换为关系模型后，系统共设计27张数据表，全部采用InnoDB引擎与utf8mb4字符集，主键统一使用BIGINT，并保留创建时间、更新时间等公共字段。核心数据表结构如表4-1至表4-7所示，其余表结构与之类似，限于篇幅不再一一列出。",
]:
    body(doc, t)

db_table(doc, '表4-1  用户账号表（user_account）', [
    ('字段名', '类型', '说明'),
    ('user_id', 'BIGINT', '用户账号ID，主键'),
    ('username', 'VARCHAR(50)', '登录名，唯一'),
    ('password_hash', 'VARCHAR(255)', 'BCrypt密码密文'),
    ('role_code', 'VARCHAR(20)', '角色：PATIENT/THERAPIST/ADMIN'),
    ('phone', 'VARCHAR(20)', '手机号，唯一'),
    ('status', 'TINYINT', '账号状态：1启用，0禁用'),
    ('last_login_at', 'DATETIME', '最后登录时间'),
    ('created_at / updated_at', 'DATETIME', '创建时间 / 更新时间'),
    ('deleted', 'TINYINT', '逻辑删除标记'),
])

db_table(doc, '表4-2  患者治疗师关系表（patient_therapist）', [
    ('字段名', '类型', '说明'),
    ('binding_id', 'BIGINT', '绑定记录ID，主键'),
    ('patient_id', 'BIGINT', '患者用户ID，外键'),
    ('therapist_id', 'BIGINT', '治疗师用户ID，外键'),
    ('relation_type', 'VARCHAR(20)', '关系类型：主责 / 协助'),
    ('status', 'VARCHAR(20)', '绑定状态：生效 / 已解除'),
    ('start_at / end_at', 'DATETIME', '生效时间 / 结束时间'),
    ('created_by', 'BIGINT', '创建人ID，外键'),
])

db_table(doc, '表4-3  动作版本表（action_version）', [
    ('字段名', '类型', '说明'),
    ('action_version_id', 'BIGINT', '动作版本ID，主键'),
    ('action_id', 'BIGINT', '所属标准动作ID'),
    ('version_no', 'INT', '版本号，与动作ID联合唯一'),
    ('action_type', 'VARCHAR(20)', '动作类型：次数型 / 保持型'),
    ('algo_config', 'JSON', '算法参数快照：置信度阈值、滤波参数、评分权重'),
    ('motion_signature', 'JSON', '动作签名，用于校验动作执行顺序'),
    ('pose_engine_id', 'VARCHAR(150)', '姿态引擎版本标识'),
    ('status', 'VARCHAR(20)', '状态：草稿 / 已生效 / 已停用'),
])

db_table(doc, '表4-4  康复计划表（rehab_plan）', [
    ('字段名', '类型', '说明'),
    ('plan_id', 'BIGINT', '康复计划ID，主键'),
    ('patient_id / therapist_id', 'BIGINT', '患者用户ID / 治疗师用户ID'),
    ('binding_id', 'BIGINT', '绑定ID，复合外键保证与绑定记录一致'),
    ('name', 'VARCHAR(100)', '计划名称，与患者、版本号联合唯一'),
    ('rehab_part_code', 'VARCHAR(30)', '康复部位编码'),
    ('version_no', 'INT', '计划版本号'),
    ('start_date / end_date', 'DATE', '计划开始日期 / 结束日期'),
    ('frequency_type / frequency_value', 'VARCHAR/INT', '训练频率类型与频率值'),
    ('status', 'VARCHAR(20)', '状态：草稿 / 已发布 / 暂停 / 结束 / 作废'),
])

db_table(doc, '表4-5  每日训练任务表（plan_daily_task）', [
    ('字段名', '类型', '说明'),
    ('task_id', 'BIGINT', '任务ID，主键'),
    ('plan_item_id', 'BIGINT', '所属计划动作项ID，外键'),
    ('task_date', 'DATE', '任务日期，与动作项联合唯一'),
    ('target_sets / target_reps', 'INT', '目标组数 / 目标次数'),
    ('status', 'VARCHAR(20)', '状态：待训练 / 进行中 / 已完成 / 已跳过'),
    ('skip_reason', 'VARCHAR(255)', '跳过原因'),
    ('started_at / completed_at', 'DATETIME', '训练开始时间 / 完成时间'),
])

db_table(doc, '表4-6  训练记录表（training_record）', [
    ('字段名', '类型', '说明'),
    ('record_id', 'BIGINT', '训练记录ID，主键'),
    ('daily_task_id', 'BIGINT', '所属每日任务ID，外键'),
    ('action_version_id', 'BIGINT', '动作版本快照，提交时复制'),
    ('started_at / ended_at', 'DATETIME', '训练开始 / 结束时间'),
    ('total_reps / correct_reps', 'INT', '总次数 / 达标次数'),
    ('total_score', 'DECIMAL(5,2)', '训练总评分，0～100'),
    ('score_detail', 'JSON', '六维分项评分快照'),
    ('avg_rom / hold_ratio', 'DECIMAL', '平均关节活动度 / 保持比例'),
    ('template_similarity', 'DECIMAL(5,4)', '模板相似度'),
    ('valid_frame_ratio', 'DECIMAL(5,4)', '有效帧比例'),
    ('pose_quality', 'VARCHAR(20)', '本次训练姿态质量评价'),
    ('client_request_id', 'VARCHAR(100)', '客户端幂等请求ID'),
])

db_table(doc, '表4-7  评估报告表（assessment_report）', [
    ('字段名', '类型', '说明'),
    ('report_id', 'BIGINT', '评估报告ID，主键'),
    ('plan_id', 'BIGINT', '所属康复计划ID，外键'),
    ('period_start / period_end', 'DATE', '报告周期起止日期，与计划联合唯一'),
    ('data_snapshot', 'JSON', '周期内训练数据快照与统计结果'),
    ('ai_draft', 'TEXT', 'AI辅助生成的报告草稿'),
    ('therapist_conclusion', 'TEXT', '治疗师审核结论与康复建议'),
    ('status', 'VARCHAR(20)', '状态：草稿 / 待审核 / 已生效'),
    ('reviewed_by / reviewed_at', 'BIGINT/DATETIME', '审核人ID / 审核时间'),
])

for t in [
    "表间一致性主要通过约束保证：康复计划通过（binding_id, patient_id, therapist_id）复合外键与绑定表保持一致，防止计划挂在已解除的绑定关系上；训练记录通过（每日任务ID, 客户端请求ID）唯一键实现重复提交的幂等处理；动作版本通过（动作ID, 版本号）唯一键保证版本号连续不重复。此外，针对常用查询建立了相应索引，如训练记录按任务与时间、按版本与时间的组合索引，公共内容按状态、部位与有效期的组合索引，以满足统计与筛选场景的查询性能要求。",
]:
    body(doc, t)

doc.save(DOC)
print('saved:', DOC)
print('paragraphs:', len(doc.paragraphs), 'tables:', len(doc.tables), 'sections:', len(doc.sections))
